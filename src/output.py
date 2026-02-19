"""
Output module: schrijft resultaten naar CSV en Excel.
"""

import logging
from pathlib import Path

import pandas as pd

logger = logging.getLogger(__name__)

# Kolomvolgorde voor export
EXPORT_COLUMNS = [
    "bwb_id",
    "titel",
    "citeertitel",
    "soort_regeling",
    "ilt_domein",
    "taaktypes",
    "ministerie_type",
    "eu_grondslag_indicatie",
    "datum_inwerkingtreding",
    "status",
    "url",
    "gevonden_op_zoektermen",
    "manco_categorie",
    "manco_toelichting",
]

# Domeinkleur voor Excel-opmaak
DOMAIN_COLORS = {
    "Luchtvaart":                    "BDD7EE",
    "Scheepvaart":                   "DDEBF7",
    "Rail":                          "E2EFDA",
    "Wegvervoer":                    "FFF2CC",
    "Gevaarlijke stoffen":           "FCE4D6",
    "Milieu & Leefomgeving":         "D9D2E9",
    "Bouw & Producten":              "F4CCCC",
    "Onbekend / nader te bepalen":   "F2F2F2",
}

MANCO_COLORS = {
    "gedekt":               "C6EFCE",
    "ongedekte_taak":       "FFC7CE",
    "slapende_bevoegdheid": "FFEB9C",
    "verouderde_wettekst":  "BDD7EE",
}


def save_csv(df: pd.DataFrame, path: str) -> None:
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    cols = [c for c in EXPORT_COLUMNS if c in df.columns]
    df[cols].to_csv(output_path, index=False, encoding="utf-8-sig")
    logger.info("CSV opgeslagen: %s (%d rijen)", output_path, len(df))


def save_excel(df: pd.DataFrame, gap_df: pd.DataFrame, path: str) -> None:
    """
    Schrijf een Excel-bestand met twee tabs:
    - 'Inventarisatie': alle gevonden regelgeving
    - 'Manco analyse': alleen de niet-gedekte regelgeving, gesorteerd op categorie
    """
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    cols = [c for c in EXPORT_COLUMNS if c in df.columns]
    gap_cols = [c for c in EXPORT_COLUMNS if c in gap_df.columns]

    mancos = gap_df[gap_df["manco_categorie"] != "gedekt"].sort_values(
        ["manco_categorie", "ilt_domein"]
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df[cols].to_excel(writer, sheet_name="Inventarisatie", index=False)
        mancos[gap_cols].to_excel(writer, sheet_name="Manco analyse", index=False)
        _apply_excel_formatting(writer)

    logger.info("Excel opgeslagen: %s", output_path)


def _apply_excel_formatting(writer: pd.ExcelWriter) -> None:
    """Pas basisopmaak toe op beide sheets."""
    try:
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        for sheet_name in writer.sheets:
            ws = writer.sheets[sheet_name]

            # Header vet + bevroren
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(wrap_text=True)
            ws.freeze_panes = "A2"

            # Kolombreedte automatisch aanpassen (max 60 tekens)
            for col_idx, col in enumerate(ws.columns, 1):
                max_len = max(
                    (len(str(cell.value or "")) for cell in col), default=10
                )
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)

            # Rijkleuring op manco_categorie (sheet 'Manco analyse')
            if sheet_name == "Manco analyse":
                header = [cell.value for cell in ws[1]]
                if "manco_categorie" in header:
                    manco_col = header.index("manco_categorie") + 1
                    for row in ws.iter_rows(min_row=2):
                        manco_val = str(row[manco_col - 1].value or "")
                        color = MANCO_COLORS.get(manco_val, "FFFFFF")
                        fill = PatternFill(fill_type="solid", fgColor=color)
                        for cell in row:
                            cell.fill = fill

    except ImportError:
        logger.warning("openpyxl styles niet beschikbaar, opmaak overgeslagen.")
